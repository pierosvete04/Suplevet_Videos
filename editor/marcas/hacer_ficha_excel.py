"""Arma marcas/ficha de marca.xlsx (la ficha en Excel de la Plantilla de marca). Se corre en la PC de desarrollo
(con openpyxl); la sala solo la copia y la lee sin librerías (importar_marca.leer_excel)."""
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

SALIDA = r"F:\TRABAJOS CLAUDE\REDES SOCIALES - SUPLEVET\1_PRODUCCION\editor\marcas\ficha de marca.xlsx"
AZUL, GRIS, LLENAR = "1F3A5F", "6B7280", "FFF7D6"
tit = Font(bold=True, color="FFFFFF", size=11)
fondo_tit = PatternFill("solid", fgColor=AZUL)
llenar = PatternFill("solid", fgColor=LLENAR)
gris = Font(color=GRIS, italic=True, size=10)
borde = Border(*(Side(style="thin", color="D1D5DB"),) * 4)
envolver = Alignment(wrap_text=True, vertical="top")

wb = Workbook()


def hoja(nombre, columnas, anchos, filas_vacias=0, ayuda=None):
    ws = wb.create_sheet(nombre)
    fila = 1
    if ayuda:
        ws.cell(1, 1, ayuda).font = Font(color=GRIS, size=10)
        ws.cell(1, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(columnas))
        ws.row_dimensions[1].height = 48
        fila = 2
    for k, (c, w) in enumerate(zip(columnas, anchos), 1):
        x = ws.cell(fila, k, c)
        x.font, x.fill, x.alignment, x.border = tit, fondo_tit, Alignment(vertical="center", wrap_text=True), borde
        ws.column_dimensions[chr(64 + k)].width = w
    ws.row_dimensions[fila].height = 30
    ws.freeze_panes = ws.cell(fila + 1, 1)
    for r in range(fila + 1, fila + 1 + filas_vacias):
        for k in range(1, len(columnas) + 1):
            x = ws.cell(r, k)
            x.fill, x.border, x.alignment = llenar, borde, envolver
    return ws, fila


# ---------- Léeme ----------
ws = wb.active
ws.title = "Léeme"
ws.column_dimensions["A"].width = 110
lineas = [
    ("FICHA DE TU MARCA", Font(bold=True, size=16, color=AZUL)),
    ("Llena las casillas amarillas de cada hoja. No cambies los títulos azules: así la sala reconoce cada respuesta.", None),
    ("Lo que no sepas, déjalo vacío: después lo completa Antigravity (o Claude) desde la sala, y tú lo revisas.", None),
    ("", None),
    ("Marca → qué vendes, a quién, cómo hablas, cómo te compran y tus tipografías.", None),
    ("Colores → cada color con su código (ej. #253C61) y para qué se usa. Si no sabes los códigos, la sala los saca de tu logo.", None),
    ("Palabras → todas las palabras que importan: el nombre de tu marca, las de tu producto, las que se resaltan en los subtítulos,", None),
    ("   las frases con las que pides la compra, lo que nunca se dice… y, si quieres, cómo la escribe mal la transcripción.", None),
    ("Imágenes y música → cómo deben verse las imágenes de tus videos y cómo debe sonar la música.", None),
    ("Públicos → a quiénes les hablas (uno por fila). Se eligen al pedir un guion.", None),
    ("", None),
    ("CUANDO TERMINES", Font(bold=True, size=12, color=AZUL)),
    ("Guarda este Excel dentro de la carpeta MI MARCA (junto a logos/, fuentes/ y productos/) y súbela en la sala:", None),
    ("Configuración → Marca → «Importar mi marca». La sala te muestra todo lo que encontró y nada se guarda hasta que lo apruebes.", None),
]
for k, (t, f) in enumerate(lineas, 1):
    ws.cell(k, 1, t).font = f or Font(size=11)

# ---------- Marca ----------
ws, f0 = hoja("Marca", ["Pregunta", "Tu respuesta", "Ejemplo"], [34, 70, 52])
preguntas = [
    ("Nombre de la marca", "Ropa Fit"),
    ("Qué vendes", "Ropa deportiva que respira y se seca rápido"),
    ("Qué te hace distinto", "Telas peruanas, tallas reales, envío en 24 horas"),
    ("A quién le vendes", "Mujeres de 20 a 35 años en Lima que entrenan 3 veces por semana"),
    ("Tono (cómo hablas)", "Cercana, con energía, palabras simples, sin tecnicismos"),
    ("Cómo te compran", "Por WhatsApp o en la web; envío a todo el Perú"),
    ("Llamado a la acción: título", "Pide la tuya"),
    ("Llamado a la acción: texto", "Envíos a todo el Perú"),
    ("Llamado a la acción: botón", "Escríbenos"),
    ("Tipografía de los títulos", "Bebas Neue"),
    ("Tipografía del texto", "DM Sans"),
    ("Tipografía de los subtítulos de los videos", "Montserrat Black"),
    ("Lo que nunca se promete", "Bajar de peso, resultados garantizados"),
    ("Más para la IA", "Siempre decimos «tu ropa de entrenar», nunca «outfit»"),
]
for k, (q, e) in enumerate(preguntas, f0 + 1):
    ws.cell(k, 1, q).font = Font(bold=True)
    ws.cell(k, 1).alignment = envolver
    b = ws.cell(k, 2)
    b.fill, b.border, b.alignment = llenar, borde, envolver
    ws.cell(k, 3, e).font = gris
    ws.cell(k, 3).alignment = envolver
    ws.row_dimensions[k].height = 32

# ---------- Colores ----------
ws, f0 = hoja("Colores", ["Nombre del color", "Código (hex)", "Para qué se usa"], [28, 18, 26], 12,
              "Un color por fila. El código empieza con # (ej. #EA8C43). «Para qué se usa»: elige de la lista. El de acento es el que resalta (las cajas de los subtítulos, los botones).")
usos = DataValidation(type="list", formula1='"principal,secundario,acento,fondo,texto,otro"', allow_blank=True)
ws.add_data_validation(usos)
usos.add(f"C{f0 + 1}:C{f0 + 12}")
for k, (n, h, u) in enumerate([("Ej.: Azul", "#253C61", "principal")], f0 + 13):
    for c, v in enumerate((n, h, u), 1):
        ws.cell(k, c, v).font = gris

# ---------- Palabras ----------
TIPOS = ["Nombre de la marca", "Palabra del producto", "Resaltar en subtítulos", "Frase para pedir la compra",
         "Nunca se dice", "Siempre tiene que ir", "Frase de cierre"]
ws, f0 = hoja("Palabras", ["Palabra o frase", "Qué es", "Cómo la escribe mal la transcripción (opcional, separadas por coma)"],
              [36, 28, 60], 60,
              "Una palabra o frase por fila y elige «Qué es» de la lista. Ej.: «Ropa Fit» → Nombre de la marca (y mal: ropa fid, ropafit) · «leggings» → Palabra del producto · «escríbenos por WhatsApp» → Frase para pedir la compra.")
dv = DataValidation(type="list", formula1='"' + ",".join(TIPOS) + '"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"B{f0 + 1}:B{f0 + 60}")

# ---------- Imágenes y música ----------
ws, f0 = hoja("Imágenes y música", ["Regla", "Para"], [90, 26], 20,
              "Una regla por fila. Ej.: «Siempre se ve una mujer entrenando con la ropa puesta» → Imágenes · «Nunca gente de espaldas» → Imágenes · «Siempre enérgica, con ritmo» → Música. «Lo que más se puede mostrar» va una sola vez.")
dv2 = DataValidation(type="list", formula1='"Imágenes,Música,Lo que más se puede mostrar"', allow_blank=True)
ws.add_data_validation(dv2)
dv2.add(f"B{f0 + 1}:B{f0 + 20}")

# ---------- Públicos ----------
ws, f0 = hoja("Públicos", ["Público"], [70], 8, "A quiénes les hablas, uno por fila (el primero es el principal). Ej.: Clientas nuevas · Clientas que ya compraron.")

wb.save(SALIDA)
print("listo:", SALIDA)
